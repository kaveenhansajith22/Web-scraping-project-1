
import csv
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

INPUT_FILE = "books.csv"
OUTPUT_FILE = "books.xlsx"

# Step 1: Read the CSV file
with open(INPUT_FILE, "r", newline="", encoding="utf-8-sig") as file:
    reader = csv.DictReader(file)
    books = list(reader)

workbook = Workbook()
sheet = workbook.active
sheet.title = "Books"

headers = ["Title", "Price (£)", "Rating", "Availability", "Book URL"]
sheet.append(headers)

exported = 0
invalid_prices = []

# Step 2: Convert prices and export records
for book in books:
    raw_price = book["price"].strip()

    # Handle common currency-symbol encoding issues
    cleaned_price = raw_price.replace("£", "").replace("Â", "").strip()

    try:
        price = Decimal(cleaned_price)
    except InvalidOperation:
        invalid_prices.append((book["title"], raw_price))
        print(f"Invalid price: {book['title']} -> {raw_price!r}")
        continue

    sheet.append([
        book["title"],
        float(price),
        book["rating"],
        book["availability"],
        book["url"]
    ])
    exported += 1

# Step 3: Format the header
for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="244062")
    cell.alignment = Alignment(vertical="center")

# Step 4: Format prices and links
for row in range(2, sheet.max_row + 1):
    sheet.cell(row, 2).number_format = '£#,##0.00'

    url_cell = sheet.cell(row, 5)
    url_cell.hyperlink = url_cell.value
    url_cell.style = "Hyperlink"

# Step 5: Set column widths
widths = {"A": 45, "B": 14, "C": 12, "D": 20, "E": 65}

for column, width in widths.items():
    sheet.column_dimensions[column].width = width

sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
sheet.row_dimensions[1].height = 25

# Step 6: Save the workbook
workbook.save(OUTPUT_FILE)

print("\n--- EXPORT REPORT ---")
print(f"Records in CSV: {len(books)}")
print(f"Records exported: {exported}")
print(f"Records with invalid prices: {len(invalid_prices)}")
print(f"Excel file saved as: {OUTPUT_FILE}")