
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT_FILE = "detailed_books.csv"
OUTPUT_FILE = "detailed_books.xlsx"

# Read the scraped CSV data
with open(INPUT_FILE, "r", encoding="utf-8-sig", newline="") as file:
    reader = csv.DictReader(file)
    headers = reader.fieldnames
    rows = list(reader)

# Create an Excel workbook
workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Detailed Books"

# Add the column headers
worksheet.append(headers)

# Add all book records
for row in rows:
    worksheet.append([row.get(header, "") for header in headers])

# Style the header row
for cell in worksheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="2C3E50"
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

# Make product URLs clickable
url_column = headers.index("url") + 1 if "url" in headers else None

if url_column:
    for row_number in range(2, worksheet.max_row + 1):
        cell = worksheet.cell(row=row_number, column=url_column)

        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

# Adjust column widths
for column_number, header in enumerate(headers, start=1):
    column_letter = get_column_letter(column_number)

    if header == "description":
        worksheet.column_dimensions[column_letter].width = 60
    elif header == "url":
        worksheet.column_dimensions[column_letter].width = 55
    elif header == "title":
        worksheet.column_dimensions[column_letter].width = 45
    else:
        worksheet.column_dimensions[column_letter].width = 22

# Keep the headers visible when scrolling
worksheet.freeze_panes = "A2"

# Add Excel filters
worksheet.auto_filter.ref = worksheet.dimensions

# Wrap long descriptions
if "description" in headers:
    description_column = headers.index("description") + 1

    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(
            row=row_number,
            column=description_column
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

# Save the Excel file
workbook.save(OUTPUT_FILE)

print("Excel export completed successfully!")
print(f"Records exported: {len(rows)}")
print(f"Excel file: {OUTPUT_FILE}")